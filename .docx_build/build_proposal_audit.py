from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"D:\ESD")
OUT = ROOT / "deliverables" / "ClawGrid_Proposal_Implementation_Audit.docx"

NAVY = "17365D"
PALE_BLUE = "EAF2F8"
LIGHT_GRAY = "D9D9D9"
GREEN = "E2F0D9"
AMBER = "FFF2CC"
RED = "FCE4D6"
BLUE = "DDEBF7"
WHITE = "FFFFFF"
BLACK = "000000"


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=90, bottom=90, end=90):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_borders(cell, color=LIGHT_GRAY, size="5"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        node = borders.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:color"), color)


def set_repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_width(cell, inches):
    cell.width = Inches(inches)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def set_font(run, name="Arial", size=10, bold=False, color=BLACK):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)


def suppress_paragraph_border(element):
    p_pr = element.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = p_bdr.find(qn("w:bottom"))
    if bottom is None:
        bottom = OxmlElement("w:bottom")
        p_bdr.append(bottom)
    bottom.set(qn("w:val"), "nil")


def set_cell_text(cell, text, *, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT, color=BLACK, size=8.5):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.05
    run = p.add_run(str(text))
    set_font(run, size=size, bold=bold, color=color)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    set_cell_margins(cell)
    set_cell_borders(cell)


def status_fill(status):
    s = status.lower()
    if s.startswith("yes"):
        return GREEN
    if s.startswith("partial"):
        return AMBER
    if s.startswith("changed"):
        return BLUE
    return RED


def add_table(doc, headers, rows, widths, status_col=None, font_size=8.5):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = "Table Grid"
    hdr = table.rows[0]
    set_repeat_header(hdr)
    for i, (head, width) in enumerate(zip(headers, widths)):
        set_width(hdr.cells[i], width)
        set_cell_text(hdr.cells[i], head, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, color=WHITE, size=8.5)
        shade(hdr.cells[i], NAVY)
    for idx, row_data in enumerate(rows):
        row = table.add_row()
        prevent_row_split(row)
        for i, (value, width) in enumerate(zip(row_data, widths)):
            set_width(row.cells[i], width)
            align = WD_ALIGN_PARAGRAPH.CENTER if i == 0 or i == status_col else WD_ALIGN_PARAGRAPH.LEFT
            set_cell_text(row.cells[i], value, align=align, size=font_size)
            if status_col is not None and i == status_col:
                shade(row.cells[i], status_fill(str(value)))
                for run in row.cells[i].paragraphs[0].runs:
                    run.font.bold = True
            elif idx % 2 == 1:
                shade(row.cells[i], PALE_BLUE)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(10 if level == 1 else 7)
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(text)
    return p


def add_body(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.12
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        set_font(r1, size=10.2, bold=True)
        r2 = p.add_run(text[len(bold_lead):])
        set_font(r2, size=10.2)
    else:
        r = p.add_run(text)
        set_font(r, size=10.2)
    return p


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.58)
section.bottom_margin = Inches(0.58)
section.left_margin = Inches(0.58)
section.right_margin = Inches(0.58)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Arial"
normal.font.size = Pt(10.2)
normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")

title_style = styles["Title"]
title_style.font.name = "Arial"
title_style.font.size = Pt(22)
title_style.font.bold = True
title_style.font.color.rgb = RGBColor(0, 0, 0)
title_style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
title_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
suppress_paragraph_border(title_style._element)

for style_name, size in (("Heading 1", 14), ("Heading 2", 11.5)):
    style = styles[style_name]
    style.font.name = "Arial"
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor(0, 0, 0)
    style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")

title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_after = Pt(7)
suppress_paragraph_border(title._element)
run = title.add_run("ClawGrid Proposal Implementation and Work Completed Report")
set_font(run, size=22, bold=True)

subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle.paragraph_format.space_after = Pt(2)
r = subtitle.add_run("Comparison with FINAL MITHUN 1")
set_font(r, size=11, bold=True, color="404040")

date_p = doc.add_paragraph()
date_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
date_p.paragraph_format.space_after = Pt(14)
r = date_p.add_run("Status date 20 October 2026")
set_font(r, size=9.5, color="606060")

add_body(doc, "This report compares every major function, hardware item, and software tool proposed in the submitted ClawGrid document with the work that has actually been executed. The central embedded platform is substantially implemented: native ESP-IDF and FreeRTOS firmware, deterministic safety control, authenticated two-board communication, a four-room virtual home, reproducible tests, and measured resource telemetry. The proposal's physical multi-room installation, metering, backup power, three-way electrical feedback, continuous microphone path, and production security remain unimplemented or only partly implemented.")
add_body(doc, "The prototype is deliberately low voltage. Physical outputs are inhibited, no sensor or relay is wired, Node A is currently unavailable, and Node B is available on COM7. Historical two-board results are identified separately from the current single-board results. A software acknowledgement is never described as proof that a physical relay or mains load changed state.")

add_heading(doc, "Status Legend", 1)
legend_rows = [
    ("Yes", "Implemented in the current project and supported by code or retained test evidence"),
    ("Partial", "Some required behavior exists, but an important function or validation step remains"),
    ("Changed", "The goal is covered through a different architecture from the submitted proposal"),
    ("No", "The proposed function or component has not been implemented"),
]
add_table(doc, ["Status", "Meaning"], legend_rows, [1.15, 6.15], status_col=0, font_size=9)

add_heading(doc, "Current Project Position", 1)
summary_rows = [
    ("Embedded platform", "Yes", "Two ESP32-S3 firmware roles use native ESP-IDF and FreeRTOS with static tasks, queues, watchdogs, core affinity, and telemetry."),
    ("Home automation", "Partial", "Four virtual rooms and eight virtual light and fan devices are controlled by MCU policy. No physical room installation is wired."),
    ("Distributed operation", "Partial", "Node A control and Node B speech/API roles were tested together historically. The latest A image is compiled but not flashed because A is unavailable."),
    ("Offline voice", "Partial", "MultiNet6 is loaded on Node B and raw PCM reaches the board with verified checksums, but the current 16-command synthetic corpus produced zero detections."),
    ("Electrical safety evidence", "No", "No auxiliary contacts, current or voltage measurements, low-voltage driver validation, or physical actuation campaign exists."),
    ("Research evidence", "Partial", "Test matrices, raw traces, provenance labels, ablations, and a novelty plan exist. Publication claims still need current paired tests, real acoustic data, physical feedback, and controlled baselines."),
]
add_table(doc, ["Area", "Implemented", "Current position"], summary_rows, [1.45, 1.0, 4.85], status_col=1, font_size=8.8)

doc.add_page_break()
add_heading(doc, "Proposal Functions Compared with Executed Work", 1)
function_rows = [
    ("1", "Switchless room-aware light and fan control", "Partial", "A controls four virtual rooms and eight virtual light and fan devices through SafetyKernel. Physical GPIOs remain inhibited and no loads are connected."),
    ("2", "One safety gateway and one controller per room", "Changed", "The prototype uses two ESP32-S3 boards: A for deterministic control and B for speech, API, and communication. Four physical room controllers are absent."),
    ("3", "Offline fixed voice commands on room nodes", "Partial", "B runs ESP-SR 1.7 MultiNet6 with 16 phrases and browser push-to-talk PCM input. Model readiness passes, but current recognition fails at 0 of 16 synthetic command matches."),
    ("4", "PIR, mmWave, and doorway transition fusion", "Partial", "A implements synthetic per-room PIR, presence, temperature, lux, and door inputs. One real PIR is owned but unwired. mmWave and reed sensors are absent."),
    ("5", "Room localization", "Changed", "The virtual house moves a controllable person between rooms and synthesizes room inputs. This is a demonstration input generator, not physical multi-room localization."),
    ("6", "ESP-Claw gateway for complex natural language", "No", "No ESP-Claw, Raspberry Pi, local LLM, llama.cpp, Ollama, or unrestricted natural-language policy generator is implemented."),
    ("7", "Restricted machine-verifiable policies", "Partial", "Scenes, timers, holds, grace settings, device intents, and bounded command schemas are validated. General JSON policies generated from complex natural language are absent."),
    ("8", "Deterministic safety kernel before actuation", "Yes", "All manual, automation, and voice intents are proposed to SafetyKernel. It enforces device permissions, minimum dwell, manual hold, and maximum run constraints."),
    ("9", "Home Away Sleep and power-limit behavior", "Partial", "Home, Away, Sleep, and All-off scenes are implemented. Measured whole-house power limits and budget enforcement are absent."),
    ("10", "Cooperative load shedding and restoration", "No", "There is no distributed load-offer protocol, measured power budget, priority negotiation, or jittered restoration controller."),
    ("11", "Three-way switching verification", "Partial", "Command identity, authenticated delivery, exact ACK payload, and a shadow evidence monitor exist. Auxiliary-contact and voltage/current feedback are not connected, so physical verification is not achieved."),
    ("12", "Fail-operational split power architecture", "No", "No always-on control branch, smart-load branch, certified panel design, or power-loss autonomy has been built."),
    ("13", "Battery-backed secondary control bus", "No", "The boards are USB powered. No LiFePO4 UPS, backed router, failover supply, or outage-duration evidence exists."),
    ("14", "Per-house encryption and secure local bus", "Yes", "ESP-NOW uses encrypted paired peers plus mandatory application HMAC, pinned identities, boot identity, ordering, issue time, and a 3-second command TTL."),
    ("15", "Reliable command confirmation", "Yes", "A 24-byte command receives a 40-byte exact-command ACK, with bounded retries and cached duplicate outcomes. This verifies software processing only."),
    ("16", "Tamper-evident hash-chained traffic log", "No", "A bounded volatile command audit exists, but there is no persistent hash chain or externally anchored audit log."),
    ("17", "Persistent policies and recovery", "Partial", "Versioned CRC-protected NVS policy records and a persistent boot counter are implemented. ON outputs, timers, and manual holds are intentionally not restored after reboot. Final policy restoration verification was interrupted."),
    ("18", "Room naming scenes holds and manual control", "Yes", "The browser interface and MCU API support four named rooms, per-device control, scenes, timers, holds, release, diagnostics, and state export."),
    ("19", "Dashboard and diagnostics", "Changed", "A responsive browser-based virtual house replaces the proposed Flutter or Android application. It works from a standard browser and does not modify the phone."),
    ("20", "Local operation without cloud dependency", "Yes", "Control, safety decisions, radio messages, and speech inference are designed for local execution. The laptop browser bridge transports data but does not perform recognition or control policy."),
    ("21", "OTA and rollback", "Partial", "Dual application slots and pending-image startup health handling exist, but there is no signed OTA transport, release service, or interrupted-update campaign."),
    ("22", "Production device security", "Partial", "Private pair credentials and authenticated traffic exist. Secure boot, flash encryption, NVS encryption, key rotation, eFuse provisioning, and rollback policy are not activated."),
    ("23", "Energy measurement and history", "Partial", "Software on-time counters and a validated DC trace integrator exist. No physical voltage/current trace, calibrated meter, kWh result, or energy-saving claim exists."),
    ("24", "Fail-safe manual intervention", "Partial", "Software All-off and maintenance-style output inhibition exist. A hardwired panel emergency OFF that bypasses firmware has not been built."),
    ("25", "High availability and self-sustaining operation", "Partial", "Watchdog recovery, OFF-by-default reboot behavior, authenticated retries, and bounded state exist. Long runs, supply interruption, network failover, and backup-power testing remain."),
]
add_table(doc, ["No", "Proposed item", "Status", "Executed work and limitation"], function_rows, [0.35, 2.05, 0.85, 4.05], status_col=2, font_size=8.1)

add_heading(doc, "Proposed Hardware Compared with Available Hardware", 1)
hardware_rows = [
    ("Safety gateway ESP32-S3", "1", "Partial", "Node A historically performs deterministic control. Its latest firmware is compiled but A is currently unavailable and has no panel installation."),
    ("Optional Raspberry Pi 5 or mini PC", "1", "No", "No edge LLM host is part of the executed architecture. The laptop provides the browser and serial bridge only."),
    ("Room ESP32-S3 controllers", "4", "No", "The project has two ESP32-S3 roles in total, not four room controllers."),
    ("INMP441 I2S microphones", "4", "No", "No I2S microphone is connected. Browser microphone capture supplies PCM for push-to-talk testing."),
    ("PIR sensors", "4", "Partial", "One PIR is available but unwired. Synthetic PIR values drive the virtual demonstration."),
    ("LD2410-class mmWave sensors", "4", "No", "No mmWave hardware or driver is present."),
    ("Door reed or transition sensors", "4", "No", "Door state is simulated in the virtual environment only."),
    ("Certified panel energy meter", "1", "No", "No isolated Modbus meter or whole-house electrical trace exists."),
    ("Branch meters", "Per branch", "No", "No calibrated branch voltage, current, or power feedback exists."),
    ("DIN contactors with auxiliary contacts", "Per channel", "No", "Two low-voltage relay modules are available but unwired. They are not treated as DIN contactors or electrical state feedback."),
    ("Optocoupler driver interfaces", "Per coil", "No", "No driver or isolation interface has been assembled or tested."),
    ("Isolated control PSU and DC rails", "1 set", "No", "Boards are powered by USB during development."),
    ("LiFePO4 UPS or protected mini UPS", "1", "No", "No backup supply or runtime-under-outage measurement exists."),
    ("Backed local router", "1", "No", "No project-specific backed access point or failover test exists."),
    ("Panel protection and emergency OFF", "As designed", "No", "No mains or panel work is permitted in this prototype phase. Software output inhibition is not a hardwired emergency stop."),
    ("DIN enclosure terminals fuses and earthing", "As designed", "No", "No enclosure or mains installation has been built."),
]
add_table(doc, ["Proposal hardware", "Quantity", "Implemented", "Actual position"], hardware_rows, [2.05, 0.75, 0.9, 3.6], status_col=2, font_size=8.2)

add_heading(doc, "Proposed Software Compared with Executed Software", 1)
software_rows = [
    ("ESP-IDF and FreeRTOS", "Yes", "Native app_main and CMake builds, static task stacks and queues, task priorities, core affinity, watchdogs, event readiness, and diagnostics are implemented."),
    ("ESP-SR WakeNet and MultiNet", "Partial", "MultiNet6 fixed-command inference is loaded. WakeNet, AFE, continuous listening, and I2S microphone capture are absent; recognition currently fails on the synthetic corpus."),
    ("ESP-NOW application framing", "Yes", "Encrypted unicast, mandatory HMAC, sequence and boot identity, expiry, exact ACK, duplicate cache, and bounded retry are implemented."),
    ("Wi-Fi and MQTT Mosquitto", "Partial", "Wi-Fi and direct local API transport are used. MQTT, retained broker state, and MQTT fault tests are absent."),
    ("FastAPI or equivalent gateway API", "Changed", "A local Python browser-to-serial bridge and an authenticated ESP32 API provide the demonstration path. It is not a production FastAPI service."),
    ("ESP-Claw with llama.cpp or Ollama", "No", "No local generative model or complex NLU runtime is present."),
    ("SQLite and Supabase", "No", "No persistent device, event, or power-history database is implemented."),
    ("Flutter or Android Studio", "Changed", "A browser-only virtual house and control dashboard replaces the native mobile application."),
    ("Secure boot flash and NVS encryption", "No", "These production security controls are not enabled."),
    ("VS Code CMake and Git", "Partial", "CMake and PlatformIO build the firmware. Reproducible scripts and hashes exist. The workspace is not initialized as a Git repository."),
    ("KiCad design", "No", "No interface PCB or isolation design is included."),
    ("pytest and MQTT test client", "Partial", "pytest-based host, integration, fault, speech, and evidence tests exist. MQTT client testing is absent because MQTT is not implemented."),
]
add_table(doc, ["Proposed software", "Implemented", "Executed work and limitation"], software_rows, [2.05, 1.0, 4.25], status_col=1, font_size=8.4)

add_heading(doc, "Additional Work Beyond the Submitted Proposal", 1)
extra_rows = [
    ("Virtual house laboratory", "Interactive four-room layout with a controllable person, eight devices, synthetic temperature, lux, door, PIR, and presence values, guided tour, dropout controls, and exportable state."),
    ("Explicit simulation boundaries", "Offline preview returns unavailable device state instead of inventing MCU decisions. Synthetic inputs, host measurements, and hardware measurements are labeled separately."),
    ("RTOS ownership model", "A single high-priority ControlTask owns deterministic state; link, diagnostics, HTTP, and speech work are separated by task and core with bounded queues and static stacks."),
    ("SafetyKernel controls", "Implemented device whitelist, minimum dwell, manual holds, resume behavior, scene handling, relative timers, maximum run cutoff, and OFF-by-default restart behavior."),
    ("Stronger command protocol", "Added boot identity, command issue time, TTL, HMAC, exact-payload acknowledgement, duplicate-result cache, and retry bounds."),
    ("NVS integrity and wear control", "Added explicit versioned records, CRC, coalesced policy writes, persistent boot counter, and a rule against restoring energized outputs."),
    ("Memory management", "Added static allocation, internal-RAM admission checks, PSRAM placement, largest-block telemetry, low-water tracking, task stack headroom, queue drop counters, and build-size reporting."),
    ("Watchdog recovery evidence", "A real task-watchdog reset, reset reason, boot increment, policy retention, OFF state, and radio recovery were recorded. The final restored-policy read still needs rerun."),
    ("Speech transport integrity", "Repaired truncated USB diagnostic responses with bounded progress and total timeouts. Added PCM CRC32 so the host and ESP32 verify identical audio bytes before evaluating recognition."),
    ("Evidence freshness monitor", "Added a 60-byte fixed-memory shadow monitor with unknown, settling, verified, and fault states; identity, freshness, sequence, lease, and distinct-domain checks; and explicit OFF recovery."),
    ("Fault and ablation matrix", "Added 4096 synthetic conditions and weak baselines to measure contract violations and conservative refusals. The model also exposes the remaining common-cause problem when both feedback channels lie."),
    ("Energy trace analysis", "Added validated trapezoidal integration for externally supplied low-voltage voltage and current traces, including provenance, units, input hashing, gap limits, and malformed-data rejection."),
    ("Reproducible evidence pack", "Added raw JSON and CSV traces, retained failures, firmware and model hashes, report generation, publication tables, evidence labels, and a controlled resume script for both boards."),
    ("Literature and novelty correction", "Reframed the research contribution around fault observability and freshness-versus-availability under concurrent MCU workloads. Broad safety-kernel and two-MCU novelty claims were removed because prior work already covers them."),
    ("Low-voltage safety boundary", "Kept all output pins inhibited until exact wiring, driver, trigger polarity, supply, and load facts are known. Mains switching is excluded from the current prototype."),
]
add_table(doc, ["Added work", "What was completed"], extra_rows, [2.15, 5.15], status_col=None, font_size=8.6)

add_heading(doc, "Completed Tests and Measured Results", 1)
test_rows = [
    ("Host regression", "HOST MEASURED", "83 tests passed", "Includes compiled actual firmware C and C++ logic, monitor cases, interface boundaries, provenance, and energy-analysis tests."),
    ("Historical paired controller", "HARDWARE MEASURED", "81 of 81 rows passed", "Earlier A and B image hashes. Current paired regression is pending because A is unavailable."),
    ("Command reliability", "HARDWARE MEASURED", "1000 of 1000 OFF outcomes", "Idempotent software outcomes with GPIO inhibited."),
    ("Command round trip", "HARDWARE MEASURED", "Median 18.789 ms, p95 44.033 ms, p99 56.754 ms", "Device ACK time only. It excludes speech recognition and physical relay movement."),
    ("HTTP loopback", "HARDWARE MEASURED", "23 of 23 passed", "Includes virtual inputs, output queries, expiry, and interface cases."),
    ("Watchdog recovery", "HARDWARE MEASURED", "11 recorded checks passed", "Real watchdog reset and recovery evidence; overall campaign remains incomplete because the final restoration read was interrupted."),
    ("Current Node B suite", "HARDWARE MEASURED", "32 of 32 rows passed", "API, transport, and shadow monitor checks; no actuation."),
    ("Evidence monitor", "SIMULATED INPUT AND HOST MEASURED", "4096 conditions", "Zero necessary-contract violations, four eligible cases conservatively refused, and two false physical assertions when both modeled feedback channels lie."),
    ("Speech transport", "SIMULATED INPUT AND HARDWARE MEASURED", "18 of 18 PCM checksums matched", "Confirms byte transport to B, not recognition quality."),
    ("Speech function", "SIMULATED INPUT AND HARDWARE MEASURED", "0 of 16 commands detected", "Current functional failure. Silence and noise caused 0 of 2 command detections; this is too small for a false-accept rate."),
    ("Physical automation", "NOT MEASURED", "Pending", "No wired PIR, relay, load, feedback contact, meter, or energy measurement."),
]
add_table(doc, ["Test area", "Evidence source", "Result", "Scope"], test_rows, [1.4, 1.55, 1.6, 2.75], status_col=None, font_size=8.0)

add_heading(doc, "Work Still Required", 1)
remaining_rows = [
    ("Restore Node A and deploy both reviewed images", "Needed for current-image paired validation, live four-room control, radio recovery, and concurrent speech-control testing."),
    ("Repair useful speech recognition", "Test a real labeled microphone recording, diagnose the decoder and acoustic front end, and compare compatible fixed-command models without tuning on the final test set."),
    ("Wire a low-voltage demonstrator", "Review exact PIR and relay pins, supply, trigger polarity, driver protection, and low-voltage load before enabling any physical output."),
    ("Add independent physical feedback", "Use auxiliary state and calibrated current or voltage evidence before claiming successful physical actuation or fault detection."),
    ("Complete publication experiments", "Run single-board and dual-board equal-workload baselines, held-out speakers and noise, RF coexistence, long-duration runs, power cycling, and independently captured energy traces."),
    ("Choose the final product scope", "Either keep the paper as a low-voltage embedded research demonstrator or separately design a certified electrical installation. The current evidence supports only the first scope."),
]
add_table(doc, ["Remaining work", "Reason and completion condition"], remaining_rows, [2.35, 4.95], status_col=None, font_size=8.8)

add_heading(doc, "Overall Conclusion", 1)
add_body(doc, "The project has progressed well beyond the submitted proposal in embedded software engineering, RTOS design, protocol integrity, memory observability, virtual demonstration, and reproducible testing. The strongest completed contribution is a deterministic two-node control platform with a realistic simulation interface and an executable evidence-freshness research framework.")
add_body(doc, "The proposal's most ambitious physical claims are still open. The current project cannot claim a four-room physical deployment, accurate offline voice assistance, cooperative energy savings, three-way electrical verification, fail-operational backup power, tamper-evident logging, or production security. Those items remain clearly separated so the final report and research paper can build on verified results without overstating the prototype.")

footer = section.footer
p = footer.paragraphs[0]
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(3)
r = p.add_run("ClawGrid proposal implementation audit  |  20 October 2026")
set_font(r, size=8, color="666666")

core_props = doc.core_properties
core_props.title = "ClawGrid Proposal Implementation and Work Completed Report"
core_props.subject = "Comparison of FINAL MITHUN 1 with executed ClawGrid work"
core_props.author = "ClawGrid Project Team"
core_props.keywords = "ClawGrid, ESP32-S3, FreeRTOS, edge AI, home automation, implementation audit"

doc.save(OUT)
print(OUT)
